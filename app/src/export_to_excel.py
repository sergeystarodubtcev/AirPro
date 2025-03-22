import pandas as pd
import matplotlib.pyplot as plt
from datetime import datetime
import os

def export_air_quality_data(dust_data, gas_data, co_data, methane_data, humidity_data, temp_data, timestamps, save_path):
    # Create a DataFrame
    df = pd.DataFrame({
        'Время': timestamps,
        'Пыль (мкг/м³)': dust_data,
        'Газ (ppm)': gas_data,
        'CO (ppm)': co_data,
        'Метан (ppm)': methane_data,
        'Влажность (%)': humidity_data,
        'Температура (°C)': temp_data
    })

    # Create Excel writer
    excel_path = save_path
    writer = pd.ExcelWriter(excel_path, engine='openpyxl')

    # Write data to Excel
    df.to_excel(writer, sheet_name='Data', index=False)

    # Access the workbook and the worksheet
    workbook = writer.book
    worksheet = writer.sheets['Data']

    # Adjust column widths
    for column in worksheet.columns:
        max_length = 0
        column_name = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        worksheet.column_dimensions[column_name].width = adjusted_width

    # Create separate charts for each parameter
    from openpyxl.chart import LineChart, Reference

    # Chart settings
    chart_height = 15
    chart_width = 15

    # Dust chart
    dust_chart = LineChart()
    dust_chart.title = "Пыль (мкг/м³)"
    dust_chart.x_axis.title = "Время"
    dust_chart.y_axis.title = "мкг/м³"
    dust_chart.height = chart_height
    dust_chart.width = chart_width
    dust_data = Reference(worksheet, min_col=2, max_col=2, min_row=1, max_row=len(df)+1)
    cats = Reference(worksheet, min_col=1, max_col=1, min_row=2, max_row=len(df)+1)
    dust_chart.add_data(dust_data, titles_from_data=True)
    dust_chart.set_categories(cats)
    worksheet.add_chart(dust_chart, "I2")

    # Gas chart
    gas_chart = LineChart()
    gas_chart.title = "Газ (ppm)"
    gas_chart.x_axis.title = "Время"
    gas_chart.y_axis.title = "ppm"
    gas_chart.height = chart_height
    gas_chart.width = chart_width
    gas_data = Reference(worksheet, min_col=3, max_col=3, min_row=1, max_row=len(df)+1)
    gas_chart.add_data(gas_data, titles_from_data=True)
    gas_chart.set_categories(cats)
    worksheet.add_chart(gas_chart, "I32")

    # CO chart
    co_chart = LineChart()
    co_chart.title = "CO (ppm)"
    co_chart.x_axis.title = "Время"
    co_chart.y_axis.title = "ppm"
    co_chart.height = chart_height
    co_chart.width = chart_width
    co_data = Reference(worksheet, min_col=4, max_col=4, min_row=1, max_row=len(df)+1)
    co_chart.add_data(co_data, titles_from_data=True)
    co_chart.set_categories(cats)
    worksheet.add_chart(co_chart, "I62")

    # Methane chart
    methane_chart = LineChart()
    methane_chart.title = "Метан (ppm)"
    methane_chart.x_axis.title = "Время"
    methane_chart.y_axis.title = "ppm"
    methane_chart.height = chart_height
    methane_chart.width = chart_width
    methane_data = Reference(worksheet, min_col=5, max_col=5, min_row=1, max_row=len(df)+1)
    methane_chart.add_data(methane_data, titles_from_data=True)
    methane_chart.set_categories(cats)
    worksheet.add_chart(methane_chart, "W2")

    # Humidity chart
    humidity_chart = LineChart()
    humidity_chart.title = "Влажность (%)"
    humidity_chart.x_axis.title = "Время"
    humidity_chart.y_axis.title = "%"
    humidity_chart.height = chart_height
    humidity_chart.width = chart_width
    humidity_data = Reference(worksheet, min_col=6, max_col=6, min_row=1, max_row=len(df)+1)
    humidity_chart.add_data(humidity_data, titles_from_data=True)
    humidity_chart.set_categories(cats)
    worksheet.add_chart(humidity_chart, "W32")

    # Temperature chart
    temp_chart = LineChart()
    temp_chart.title = "Температура (°C)"
    temp_chart.x_axis.title = "Время"
    temp_chart.y_axis.title = "°C"
    temp_chart.height = chart_height
    temp_chart.width = chart_width
    temp_data = Reference(worksheet, min_col=7, max_col=7, min_row=1, max_row=len(df)+1)
    temp_chart.add_data(temp_data, titles_from_data=True)
    temp_chart.set_categories(cats)
    worksheet.add_chart(temp_chart, "W62")

    # Save the Excel file
    writer.close()

    return excel_path

# Modify the AirQualityMonitor class to add export functionality
def add_export_button(app):
    from .air_quality_monitor2 import RoundedButton
    import customtkinter as ctk
    
    # Create export button without an image to avoid the warning
    export_button = ctk.CTkButton(
        app.buttons_frame,
        text="Экспорт в Excel",
        command=lambda: export_data(app),
        corner_radius=10,
        border_width=0,
        hover=True,
        fg_color="#0078D7",
        hover_color="#005A9E",
        text_color="#FFFFFF",
        height=36
    )
    export_button.pack(fill='x', pady=5)

def export_data(app):
    try:
        # Check if there's data to export
        if not app.dust_data:
            from tkinter import messagebox
            messagebox.showwarning("Предупреждение", "Нет данных для экспорта. Подключитесь к устройству и получите данные сначала.")
            return
            
        # Ask user where to save the file
        from tkinter import filedialog
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        default_name = f'air_quality_data_{timestamp}'
        save_path = filedialog.asksaveasfilename(
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            title="Выберите место для сохранения данных"
        )
        
        if not save_path:  # User cancelled
            return
            
        excel_path = export_air_quality_data(
            list(app.dust_data),
            list(app.gas_data),
            list(app.co_data),
            list(app.methane_data),
            list(app.humidity_data),
            list(app.temp_data),
            list(app.timestamps),
            save_path
        )
        
        app.log(f"Данные успешно экспортированы в файл: {excel_path}")
        # Add a messagebox to show success
        from tkinter import messagebox
        messagebox.showinfo("Успех", f"Данные успешно экспортированы в файл:\n{excel_path}")
    except Exception as e:
        error_msg = str(e)
        app.log(f"Ошибка при экспорте данных: {error_msg}")
        # Add error messagebox
        from tkinter import messagebox
        messagebox.showerror("Ошибка", f"Не удалось экспортировать данные:\n{error_msg}")
